"""Prints the key outputs of micron_dcf.xlsx after it has been recalculated in Excel, and checks for formula errors."""
from openpyxl import load_workbook
wb = load_workbook("micron_dcf.xlsx", data_only=True); wf = load_workbook("micron_dcf.xlsx")
errs = [(ws.title, c.coordinate, wb[ws.title][c.coordinate].value) for ws in wf.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=") and isinstance(wb[ws.title][c.coordinate].value, str)
        and wb[ws.title][c.coordinate].value.startswith("#")]
print("formula errors:", errs or "none")
s = wb["Summary"]
for r in range(5, 14):
    print(f"{s.cell(r,1).value:32s}", *[f"{v:>12,.2f}" if isinstance(v, float) and abs(v) > 2 else (f"{v:>12.1%}" if isinstance(v, float) else f"{str(v):>12s}") for v in [s.cell(r, k).value for k in range(2, 6)]])
for r in (15, 16, 17, 18): print(s.cell(r, 1).value, round(s.cell(r, 2).value, 4))
d = wb["DCF"]
for name, top in (("Bear", 3), ("Base", 37), ("Bull", 71)):
    fcf = [round(d.cell(top + 11, k).value or 0) for k in range(4, 9)]
    print(name, "FCF FY27-31:", fcf)
rv = wb["Reverse DCF"]
print("sensitivity header", [rv.cell(20, k).value for k in range(3, 10)])
for r in range(21, 29): print(rv.cell(r, 1).value, [round(rv.cell(r, k).value) for k in range(3, 10)])
