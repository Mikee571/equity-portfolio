"""Checks that models.py reproduces the Excel models (values cached by Excel in valuation/*/*.xlsx).

Run from valuation/: python -m unittest test_models -v
"""
import unittest
from pathlib import Path

from openpyxl import load_workbook

import models

ROOT = Path(__file__).resolve().parent.parent


def excel_summary(path):
    s = load_workbook(path, data_only=True)["Summary"]
    return {s.cell(r, 1).value: [s.cell(r, k).value for k in range(2, 6)] for r in range(5, 21)}


class TestAgainstExcel(unittest.TestCase):
    def test_micron(self):
        xl = excel_summary(ROOT / "valuation/micron/micron_dcf.xlsx")
        for k, name in enumerate(("Bear", "Base", "Bull")):
            v = models.micron_value(models.MICRON_SCENARIOS[name])[0]
            self.assertAlmostEqual(v, xl["Value per share ($)"][k], places=2, msg=name)
        self.assertAlmostEqual(models.micron_implied_margin(),
                               xl["Margin from FY29 needed to justify the price"][0], places=4)

    def test_gtt(self):
        xl = excel_summary(ROOT / "valuation/gtt/gtt_dcf.xlsx")
        for k, name in enumerate(("Bear", "Base", "Bull")):
            s = models.GTT_SCENARIOS[name]
            v = models.gtt_value(s["h2"], s["orders"], s["marine"], s["g"])[0]
            self.assertAlmostEqual(v, xl["Value per share (€)"][k], places=2, msg=name)
        self.assertAlmostEqual(models.gtt_implied_orders(),
                               xl["Orders a year (2027-35) needed to justify the price"][0], places=3)

    def test_jpm(self):
        xl = excel_summary(ROOT / "valuation/jpm/jpm_ri_model.xlsx")
        for k, name in enumerate(("Bear", "Base", "Bull")):
            s = models.JPM_SCENARIOS[name]
            self.assertAlmostEqual(models.jpm_value(s["rote"], s["lr"], s["payout"]), xl["Value per share ($)"][k],
                                   places=2, msg=name)
        self.assertAlmostEqual(models.jpm_implied_rote(), xl["ROTE the price implies, forever"][0], places=6)


if __name__ == "__main__":
    unittest.main()
