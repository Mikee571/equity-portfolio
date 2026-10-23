"""Builds gtt_dcf.xlsx: a three-scenario DCF for GTT (Gaztransport & Technigaz, Euronext Paris: GTT) with live formulas.

Revenue engine: royalties already in the order book (company schedule) + future orders x royalty per vessel,
recognised over construction. Sources: GTT FY2025 results press release (2026-02-19), H1 2026 results press
release (2026-07-28), Yahoo Finance (price, consensus) and Trading Economics (OAT yield), both as of mid-October 2026.
Run: python build_model.py, then open in Excel (or recalculate) to cache values.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = Path(__file__).parent
OUT = HERE / "gtt_dcf.xlsx"

BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="DDE6F0")
EUR = '#,##0.0;(#,##0.0);"-"'
EUR2 = '€#,##0.00;(€#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
I = "Inputs!"
YEARS = list(range(2026, 2036))           # model columns C..L
YC = [chr(ord("C") + i) for i in range(len(YEARS))]
COL = dict(zip(YEARS, YC))


def font(color=BLACK, bold=False, italic=False):
    return Font(name="Arial", size=10, color=color, bold=bold, italic=italic)


def put(ws, ref, value, color=BLACK, fmt=None, bold=False, fill=None, italic=False):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def header_row(ws, row, labels):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=1 + j, value=lab)
        c.font = font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(horizontal="left" if j == 0 else "center")


wb = Workbook()

# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Inputs"
inp.column_dimensions["A"].width = 50
for c in YC + ["B", "M"]:
    inp.column_dimensions[c].width = 10
inp.column_dimensions["N"].width = 90
put(inp, "A1", "GTT (Gaztransport & Technigaz): inputs and scenario levers (€mm unless stated)", bold=True)
put(inp, "A2", "Blue = hardcoded input. Yellow = key assumption. Black = formula. Sources in column N.", italic=True)

ROWS = [
    (4, "Market data and capital structure", None, None, None),
    (5, "Share price (€)", 211.60, EUR2, "Yahoo Finance, close 2026-10-16"),
    (6, "Shares outstanding excl. treasury (mm)", 37.067, "0.000", "H1 2026 press release: weighted average 37,066,945 shares"),
    (7, "Cash at 30 June 2026", 384.7, EUR, "H1 2026 press release, other consolidated financial data"),
    (8, "Debt at 30 June 2026 (Danelec acquisition loan)", 90.0, EUR, "H1 2026 press release"),
    (9, "Net cash", "=B7-B8", EUR, "Formula (294.7 in the press release)"),
    (11, "Discount rate", None, None, None),
    (12, "Risk-free rate (French OAT 10-year)", 0.0475, PCT, "Trading Economics, 2026-10-07: ~4.75% (spread over Bunds at a post-2011 high)"),
    (13, "Beta", 0.90, "0.00", "Assumption: regression beta 0.42 (Yahoo) looks too low for a business tied to shipbuilding orders"),
    (14, "Equity risk premium", 0.055, PCT, "Assumption (author), Europe"),
    (15, "Cost of equity", "=B12+B13*B14", PCT, "CAPM"),
    (16, "Pre-tax cost of debt", 0.04, PCT, "Assumption"),
    (17, "Tax rate", 0.16, PCT, "H1 2026 effective rate 15.2% (37.7 / 248.2); French IP box applies to royalties"),
    (18, "Market capitalisation", "=B5*B6", EUR, "Formula"),
    (19, "WACC", "=(B18*B15+B8*B16*(1-B17))/(B18+B8)", PCT, "Formula, market-value weights"),
    (21, "Revenue engine and costs", None, None, None),
    (22, "Royalty per LNG-carrier-equivalent order (€mm, 2026 terms)", 9.5, EUR,
     "Derived: H1 2026 orders (65 units, ~60 LNGC-equivalent) added ~€570mm to the order-book revenue schedule"),
    (23, "Royalty escalation per year", 0.02, PCT, "Assumption: shipbuilding price inflation"),
    (24, "Recognition: year after order", 0.20, PCT, "Derived from the change in the order-book schedule Dec-25 -> Jun-26: 2027 +113, 2028 +317, 2029+ +140"),
    (25, "Recognition: second year after order", 0.55, PCT, "Same derivation"),
    (26, "Recognition: third year after order", 0.25, PCT, "Same derivation"),
    (27, "  Check (must be 100%)", "=SUM(B24:B26)", PCT, "Formula"),
    (28, "Core operating costs 2026 (excl. GTT Marine)", 193.0, EUR,
     "H1 2026: revenue 387.3 - EBITDA 263.6 = 123.7 of costs; ~250/yr less GTT Marine costs (~57)"),
    (29, "Core cost growth per year", 0.03, PCT, "Assumption: wages and R&D"),
    (30, "GTT Marine revenue 2026E", 70.0, EUR, "H1 2026: 31.8 (Danelec consolidated); assumption for H2"),
    (31, "GTT Marine EBITDA margin", 0.18, PCT, "H1 2026: EBITDA 5.7 on revenue 31.8"),
    (32, "Services and other revenue 2026E", 22.0, EUR, "H1 2026 services 10.8; FY2025 23.0"),
    (33, "Services growth per year", 0.02, PCT, "Assumption"),
    (34, "D&A per year", 30.0, EUR, "H1 2026 EBITDA - EBIT = 15.4 (includes Danelec purchase-price amortisation)"),
    (35, "Capex per year (net of subsidies)", 30.0, EUR, "Assumption: H1 2026 gross 14.5 less 12.3 subsidy; 2024 68.5 incl. headquarters"),
    (36, "Change in NWC as % of change in revenue", 0.10, PCT, "Assumption; FY2025 working capital rose 26.4 on +161.7 revenue"),
    (37, "H2 2026 FCF as share of modelled FY2026 FCF", 0.50, PCT, "Cash generated after 30 June 2026, added to net cash"),
]
for r, label, val, fmt, src in ROWS:
    put(inp, f"A{r}", label, bold=val is None)
    if val is None:
        continue
    is_formula = isinstance(val, str) and val.startswith("=")
    put(inp, f"B{r}", val, BLACK if is_formula else BLUE, fmt)
    put(inp, f"N{r}", src, italic=True)
for r in (13, 22):
    inp[f"B{r}"].fill = YELLOW

put(inp, "A39", "Revenue already secured by the order book (company schedule at 30 June 2026)", bold=True)
header_row(inp, 40, ["", ""] + [str(y) for y in YEARS])
BACKLOG = {2026: 644.7, 2027: 655.0, 2028: 603.0, 2029: 221.3, 2030: 73.8}
put(inp, "A41", "  Core business order-book revenue")
for y in YEARS:
    put(inp, f"{COL[y]}41", BACKLOG.get(y, 0.0), BLUE, EUR)
put(inp, "N41", "H1 2026 press release: 1,853 = H2-26 300 + 2027 655 + 2028 603 + 2029+ 295. 2026 = H1 actual 344.7 + 300; "
                "2029+ split 75/25 over 2029-30 (assumption)", italic=True)

SCEN = {
    "Bear": dict(start=44, h2=15, orders=35, marine=0.03, g=0.00,
                 story="LNG shipping is over-ordered after the 2022-26 wave and competing membrane designs win share in China: 35 orders a year."),
    "Base": dict(start=49, h2=30, orders=60, marine=0.08, g=0.015,
                 story="Company's own 10-year estimate (c.550 LNG carriers + ethane, FSRU, FLNG, onshore) ~ 60 LNGC-equivalent a year."),
    "Bull": dict(start=54, h2=40, orders=75, marine=0.12, g=0.02,
                 story="LNG demand upgrades (US, Qatar, Mozambique) and fleet renewal push orders well above the company's estimate."),
}
put(inp, "A43", "Scenarios: new orders in LNG-carrier equivalents (2026 column = second half only)", bold=True)
for name, s in SCEN.items():
    r = s["start"]
    put(inp, f"A{r}", f"{name} case", bold=True)
    put(inp, f"N{r}", s["story"], italic=True)
    put(inp, f"A{r+1}", "  New orders (units)")
    for y in YEARS:
        put(inp, f"{COL[y]}{r+1}", s["h2"] if y == 2026 else s["orders"], BLUE, "0")
    put(inp, f"A{r+2}", "  GTT Marine revenue growth")
    put(inp, f"C{r+2}", s["marine"], BLUE, PCT)
    put(inp, f"A{r+3}", "  Terminal growth")
    put(inp, f"C{r+3}", s["g"], BLUE, PCT, fill=YELLOW)
# Reverse-DCF helper cases: base everything, orders from 2027 = 0 or 1
for name, start, q in (("Rev0", 59, 0), ("Rev1", 64, 1)):
    put(inp, f"A{start}", f"Reverse-DCF helper ({'orders = 0' if q == 0 else 'orders = 1'} from 2027; rest = base)", bold=True)
    put(inp, f"A{start+1}", "  New orders (units)")
    for y in YEARS:
        put(inp, f"{COL[y]}{start+1}", "=C50" if y == 2026 else q, GREEN if y == 2026 else BLUE, "0")
    put(inp, f"A{start+2}", "  GTT Marine revenue growth")
    put(inp, f"C{start+2}", "=C51", GREEN, PCT)
    put(inp, f"A{start+3}", "  Terminal growth")
    put(inp, f"C{start+3}", "=C52", GREEN, PCT)
    SCEN[name] = dict(start=start)

put(inp, "A69", "Scenario probabilities", bold=True)
for k, (name, p) in enumerate([("Bear", 0.25), ("Base", 0.50), ("Bull", 0.25)]):
    put(inp, f"A{70+k}", f"  {name}")
    put(inp, f"B{70+k}", p, BLUE, PCT, fill=YELLOW)
put(inp, "N70", "Author's judgement", italic=True)
put(inp, "A73", "  Check (must be 100%)")
put(inp, "B73", "=SUM(B70:B72)", fmt=PCT)

put(inp, "A75", "For reference", bold=True)
REF = [("Company 2026 revenue guidance, low / high", 740, 780), ("Company 2026 EBITDA guidance, low / high", 490, 530),
       ("Consensus revenue 2026 / 2027", 768.3, 814.8), ("Consensus EPS 2026 / 2027 (€)", 11.31, 12.11)]
for k, (lab, a, b) in enumerate(REF):
    put(inp, f"A{76+k}", f"  {lab}")
    put(inp, f"B{76+k}", a, BLUE, EUR)
    put(inp, f"C{76+k}", b, BLUE, EUR)
put(inp, "N76", "H1 2026 press release (guidance confirmed)", italic=True)
put(inp, "N78", "Yahoo Finance analyst estimates, 2026-10-16 (12 analysts; mean target €234.25)", italic=True)
put(inp, "A80", "  Dividend policy: payout ~80% of consolidated net income (FY2025 dividend €8.94; interim 2026 €4.30)", italic=True)

# ---------------------------------------------------------------- Historical
his = wb.create_sheet("Historical")
his.column_dimensions["A"].width = 44
for c in "BCDEFGH":
    his.column_dimensions[c].width = 11
put(his, "A1", "GTT: historical figures (€mm)", bold=True)
put(his, "A2", "Sources: GTT FY2025 and H1 2026 press releases; FY2021-23 from company figures via stockanalysis.com / Yahoo Finance.",
    italic=True)
header_row(his, 4, ["€mm", "2021", "2022", "2023", "2024", "2025", "H1 2025", "H1 2026"])
HIST = [("Revenue", [314.7, 307.3, 427.7, 641.4, 803.0, 388.7, 387.3]),
        ("EBITDA", [None, 163.1, 241.2, 388.1, 541.8, 264.5, 263.6]),
        ("Operating income (EBIT)", [164.6, 152.2, 232.4, 374.3, 521.3, 257.1, 248.2]),
        ("Net income", [134.1, 128.3, 201.4, 347.8, 413.6, 180.0, 210.4]),
        ("EPS (€)", [3.62, 3.46, 5.43, 9.40, 11.20, 4.86, 5.68]),
        ("Dividend per share (€)", [3.10, 3.10, 4.36, 7.50, 8.94, None, 4.30])]
for i, (lab, vals) in enumerate(HIST):
    r = 5 + i
    put(his, f"A{r}", lab, bold=lab == "Revenue")
    for j, v in enumerate(vals):
        if v is not None:
            put(his, f"{chr(ord('B') + j)}{r}", v, BLUE, EUR2 if "€" in lab else EUR)
put(his, "A12", "EBITDA margin")
put(his, "A13", "EBIT margin")
put(his, "A14", "Payout (dividend / EPS)")
for j in range(7):
    c = chr(ord("B") + j)
    put(his, f"{c}12", f'=IF({c}6="","",{c}6/{c}5)', fmt=PCT)
    put(his, f"{c}13", f"={c}7/{c}5", fmt=PCT)
    put(his, f"{c}14", f'=IF({c}10="","",{c}10/{c}9)', fmt=PCT)
put(his, "A16", "Order book (core business, excl. LNG as fuel)", bold=True)
header_row(his, 17, ["", "Units", "Revenue schedule (€mm)", "", "", "", "", ""])
put(his, "A18", "31 Dec 2025")
put(his, "B18", 288, BLUE, "0")
put(his, "C18", "2026: 609 · 2027: 542 · 2028: 286 · 2029+: 155 (total 1,592)", italic=True)
put(his, "A19", "30 Jun 2026")
put(his, "B19", 306, BLUE, "0")
put(his, "C19", "H2-26: 300 · 2027: 655 · 2028: 603 · 2029+: 295 (total 1,853)", italic=True)
put(his, "A20", "Orders: 2022 peak 162 units; 2025: 45 (37 LNG carriers); H1 2026: 65 (56 LNG carriers)", italic=True)
put(his, "A21", "Company 10-year order estimate 2026-2035: c.550 LNG carriers, 25-40 ethane carriers, ≤10 FSRUs, ≤10 FLNGs, c.30 onshore tanks",
    italic=True)

# ---------------------------------------------------------------- DCF
dcf = wb.create_sheet("DCF")
dcf.column_dimensions["A"].width = 44
dcf.column_dimensions["B"].width = 4
for c in YC + ["M"]:
    dcf.column_dimensions[c].width = 10
put(dcf, "A1", "DCF by scenario (€mm). FY = calendar year. Cash flows 2027-2035 discounted from end-2026; "
               "H2 2026 cash flow added to net cash.", bold=True)
FLOW = ["Order-book revenue (already signed)", "New orders (units)", "Royalty per unit", "Value of new orders",
        "Revenue from new orders", "Core containment revenue", "GTT Marine revenue", "Services and other",
        "Total revenue", "Core operating costs", "GTT Marine costs", "EBITDA", "EBITDA margin", "D&A", "EBIT",
        "Taxes", "NOPAT", "Capex", "Change in NWC", "Free cash flow", "Discount period (years)", "Discount factor",
        "PV of FCF"]
VAL = ["Terminal growth", "Terminal value at 2035", "PV of terminal value", "Sum of PV of FCF (2027-35)",
       "Enterprise value", "+ Net cash at 30 June 2026", "+ H2 2026 free cash flow", "Equity value",
       "Value per share (€)", "Upside / (downside) vs price", "Terminal value as % of EV",
       "Implied P/E on 2027 NOPAT", "Average new orders 2027-35 (units/yr)"]
BLOCK = {}


def block(name, top):
    sr = SCEN[name]["start"]
    put(dcf, f"A{top}", f"{name} case", bold=True)
    header_row(dcf, top + 1, ["€mm", ""] + [str(y) for y in YEARS])
    rr = {lab: top + 2 + i for i, lab in enumerate(FLOW)}
    for lab, r in rr.items():
        put(dcf, f"A{r}", lab, bold=lab in ("Total revenue", "EBITDA", "Free cash flow"))
    p1, p2, p3 = f"{I}$B$24", f"{I}$B$25", f"{I}$B$26"
    for k, y in enumerate(YEARS):
        c = COL[y]
        prev = YC[k - 1] if k else None
        put(dcf, f"{c}{rr['Order-book revenue (already signed)']}", f"={I}{c}$41", GREEN, EUR)
        put(dcf, f"{c}{rr['New orders (units)']}", f"={I}{c}{sr+1}", GREEN, "0")
        put(dcf, f"{c}{rr['Royalty per unit']}", f"={I}$B$22*(1+{I}$B$23)^({y}-2026)", fmt=EUR)
        put(dcf, f"{c}{rr['Value of new orders']}",
            f"={c}{rr['New orders (units)']}*{c}{rr['Royalty per unit']}", fmt=EUR)
        v = rr["Value of new orders"]
        terms = [f"{YC[k - d]}{v}*{p}" for d, p in ((1, p1), (2, p2), (3, p3)) if k - d >= 0]
        put(dcf, f"{c}{rr['Revenue from new orders']}", "=" + ("+".join(terms) if terms else "0"), fmt=EUR)
        put(dcf, f"{c}{rr['Core containment revenue']}",
            f"={c}{rr['Order-book revenue (already signed)']}+{c}{rr['Revenue from new orders']}", fmt=EUR)
        put(dcf, f"{c}{rr['GTT Marine revenue']}",
            f"={I}$B$30" if k == 0 else f"={prev}{rr['GTT Marine revenue']}*(1+{I}$C${sr+2})",
            GREEN if k == 0 else BLACK, EUR)
        put(dcf, f"{c}{rr['Services and other']}", f"={I}$B$32*(1+{I}$B$33)^({y}-2026)", fmt=EUR)
        put(dcf, f"{c}{rr['Total revenue']}",
            f"={c}{rr['Core containment revenue']}+{c}{rr['GTT Marine revenue']}+{c}{rr['Services and other']}",
            fmt=EUR, bold=True)
        put(dcf, f"{c}{rr['Core operating costs']}", f"=-{I}$B$28*(1+{I}$B$29)^({y}-2026)", fmt=EUR)
        put(dcf, f"{c}{rr['GTT Marine costs']}", f"=-{c}{rr['GTT Marine revenue']}*(1-{I}$B$31)", fmt=EUR)
        put(dcf, f"{c}{rr['EBITDA']}",
            f"={c}{rr['Total revenue']}+{c}{rr['Core operating costs']}+{c}{rr['GTT Marine costs']}", fmt=EUR, bold=True)
        put(dcf, f"{c}{rr['EBITDA margin']}", f"={c}{rr['EBITDA']}/{c}{rr['Total revenue']}", fmt=PCT)
        put(dcf, f"{c}{rr['D&A']}", f"=-{I}$B$34", GREEN, EUR)
        put(dcf, f"{c}{rr['EBIT']}", f"={c}{rr['EBITDA']}+{c}{rr['D&A']}", fmt=EUR)
        put(dcf, f"{c}{rr['Taxes']}", f"=-{c}{rr['EBIT']}*{I}$B$17", fmt=EUR)
        put(dcf, f"{c}{rr['NOPAT']}", f"={c}{rr['EBIT']}+{c}{rr['Taxes']}", fmt=EUR)
        put(dcf, f"{c}{rr['Capex']}", f"=-{I}$B$35", GREEN, EUR)
        put(dcf, f"{c}{rr['Change in NWC']}",
            "=0" if k == 0 else f"=-{I}$B$36*({c}{rr['Total revenue']}-{prev}{rr['Total revenue']})", fmt=EUR)
        put(dcf, f"{c}{rr['Free cash flow']}",
            f"={c}{rr['NOPAT']}-{c}{rr['D&A']}+{c}{rr['Capex']}+{c}{rr['Change in NWC']}", fmt=EUR, bold=True)
        if k > 0:
            put(dcf, f"{c}{rr['Discount period (years)']}", k, BLUE, "0")
            put(dcf, f"{c}{rr['Discount factor']}", f"=1/(1+{I}$B$19)^{c}{rr['Discount period (years)']}", fmt="0.000")
            put(dcf, f"{c}{rr['PV of FCF']}", f"={c}{rr['Free cash flow']}*{c}{rr['Discount factor']}", fmt=EUR)
    last, first = YC[-1], YC[0]
    v0 = top + 2 + len(FLOW) + 1
    vr = {lab: v0 + i for i, lab in enumerate(VAL)}
    for lab, r in vr.items():
        put(dcf, f"A{r}", lab, bold=lab in ("Enterprise value", "Equity value", "Value per share (€)"))
    put(dcf, f"C{vr['Terminal growth']}", f"={I}C{sr+3}", GREEN, PCT)
    put(dcf, f"C{vr['Terminal value at 2035']}",
        f"={last}{rr['Free cash flow']}*(1+C{vr['Terminal growth']})/({I}$B$19-C{vr['Terminal growth']})", fmt=EUR)
    put(dcf, f"C{vr['PV of terminal value']}",
        f"=C{vr['Terminal value at 2035']}*{last}{rr['Discount factor']}", fmt=EUR)
    put(dcf, f"C{vr['Sum of PV of FCF (2027-35)']}", f"=SUM(D{rr['PV of FCF']}:{last}{rr['PV of FCF']})", fmt=EUR)
    put(dcf, f"C{vr['Enterprise value']}",
        f"=C{vr['PV of terminal value']}+C{vr['Sum of PV of FCF (2027-35)']}", fmt=EUR, bold=True)
    put(dcf, f"C{vr['+ Net cash at 30 June 2026']}", f"={I}$B$9", GREEN, EUR)
    put(dcf, f"C{vr['+ H2 2026 free cash flow']}", f"={first}{rr['Free cash flow']}*{I}$B$37", fmt=EUR)
    put(dcf, f"C{vr['Equity value']}",
        f"=SUM(C{vr['Enterprise value']}:C{vr['+ H2 2026 free cash flow']})", fmt=EUR, bold=True)
    put(dcf, f"C{vr['Value per share (€)']}", f"=C{vr['Equity value']}/{I}$B$6", fmt=EUR2, bold=True)
    put(dcf, f"C{vr['Upside / (downside) vs price']}", f"=C{vr['Value per share (€)']}/{I}$B$5-1", fmt=PCT)
    put(dcf, f"C{vr['Terminal value as % of EV']}", f"=C{vr['PV of terminal value']}/C{vr['Enterprise value']}", fmt=PCT)
    put(dcf, f"C{vr['Implied P/E on 2027 NOPAT']}", f"={I}$B$18/D{rr['NOPAT']}", fmt=MULT)
    put(dcf, f"C{vr['Average new orders 2027-35 (units/yr)']}",
        f"=AVERAGE(D{rr['New orders (units)']}:{last}{rr['New orders (units)']})", fmt="0.0")
    BLOCK[name] = dict(rows=rr, vals=vr)
    return v0 + len(VAL) + 2


top = 3
for name in ("Bear", "Base", "Bull", "Rev0", "Rev1"):
    if name == "Rev0":
        put(dcf, f"A{top}", "Helper blocks for the reverse DCF (base case with 0 and 1 new orders a year from 2027)",
            italic=True)
        top += 2
    top = block(name, top)

# ---------------------------------------------------------------- Reverse DCF
rev = wb.create_sheet("Reverse DCF")
rev.column_dimensions["A"].width = 50
for c in "BCDEFGHIJKL":
    rev.column_dimensions[c].width = 10
put(rev, "A1", "What do you have to believe? Orders a year (LNG-carrier equivalents, 2027-35) that justify the price; "
               "everything else as in the base case.", bold=True)
put(rev, "A2", "Value is linear in orders, so it is solved exactly from two helper DCFs (0 and 1 order a year).", italic=True)
e0 = f"DCF!C{BLOCK['Rev0']['vals']['Equity value']}"
e1 = f"DCF!C{BLOCK['Rev1']['vals']['Equity value']}"
put(rev, "A4", "Target equity value (price x shares)")
put(rev, "C4", f"={I}B18", GREEN, EUR)
put(rev, "A5", "Equity value with 0 orders a year")
put(rev, "C5", f"={e0}", GREEN, EUR)
put(rev, "A6", "Extra equity value per order a year")
put(rev, "C6", f"={e1}-{e0}", fmt=EUR)
put(rev, "A7", "Implied orders a year, 2027-2035", bold=True)
put(rev, "C7", "=(C4-C5)/C6", fmt="0.0", bold=True)
put(rev, "A8", "  Base case assumption")
put(rev, "C8", "=Inputs!D50", GREEN, "0")
put(rev, "A9", "  Company 10-year estimate (c.550 LNGC + other units, per year)")
put(rev, "C9", 60, BLUE, "0")

r0, r1 = BLOCK["Rev0"]["rows"], BLOCK["Rev1"]["rows"]
put(rev, "A12", "Free cash flow rows used by the grid (€mm)", bold=True)
header_row(rev, 13, ["", ""] + [str(y) for y in YEARS[1:]])
put(rev, "A14", "F0: FCF with 0 orders a year")
put(rev, "A15", "dF: FCF per order a year")
put(rev, "A16", "Discount period")
for k, y in enumerate(YEARS[1:]):
    c = chr(ord("C") + k)
    src = COL[y]
    put(rev, f"{c}14", f"=DCF!{src}{r0['Free cash flow']}", GREEN, EUR)
    put(rev, f"{c}15", f"=DCF!{src}{r1['Free cash flow']}-DCF!{src}{r0['Free cash flow']}", fmt=EUR)
    put(rev, f"{c}16", k + 1, BLUE, "0")
G = f"{I}$C$52"
CONST = f"({I}$B$9+DCF!C{BLOCK['Rev0']['vals']['+ H2 2026 free cash flow']})"
put(rev, "A19", "Value per share (€): WACC (rows) x new orders a year 2027-35 (columns)", bold=True)
ORDERS = [30, 40, 50, 60, 70, 80, 90]
WACCS = [0.075, 0.08, 0.085, 0.09, 0.095, 0.10, 0.11]
header_row(rev, 20, ["WACC \\ orders", ""] + [""] * len(ORDERS))
for k, q in enumerate(ORDERS):
    put(rev, f"{chr(ord('C') + k)}20", q, BLUE, "0", bold=True)
for i, w in enumerate(WACCS):
    r = 21 + i
    put(rev, f"A{r}", w, BLUE, PCT, bold=True)
    for k in range(len(ORDERS)):
        c = chr(ord("C") + k)
        put(rev, f"{c}{r}",
            f"=(SUMPRODUCT(($C$14:$K$14+{c}$20*$C$15:$K$15)/(1+$A{r})^$C$16:$K$16)"
            f"+($K$14+{c}$20*$K$15)*(1+{G})/($A{r}-{G})/(1+$A{r})^$K$16+{CONST})/{I}$B$6", fmt="€#,##0")
put(rev, "A29", "Current price (€)")
put(rev, "C29", f"={I}B5", GREEN, EUR2)
put(rev, "A30", "Model WACC")
put(rev, "C30", f"={I}B19", GREEN, PCT)

# ---------------------------------------------------------------- Summary
sm = wb.create_sheet("Summary", 0)
sm.column_dimensions["A"].width = 46
for c in "BCDE":
    sm.column_dimensions[c].width = 15
put(sm, "A1", "GTT (Gaztransport & Technigaz): three-scenario DCF", bold=True)
put(sm, "A2", "Valuation date 2026-10-16. €mm unless stated. Change blue/yellow cells in 'Inputs'; everything else recalculates.",
    italic=True)
header_row(sm, 4, ["", "Bear", "Base", "Bull", "Prob.-weighted"])


def link(n, key, kind="vals", col="C"):
    return f"=DCF!{col}{BLOCK[n][kind][key]}"


SUMMARY = [("New orders a year 2027-35 (LNGC-eq.)", lambda n: link(n, "Average new orders 2027-35 (units/yr)"), "0"),
           ("2026 revenue", lambda n: link(n, "Total revenue", "rows", "C"), EUR),
           ("2030 revenue", lambda n: link(n, "Total revenue", "rows", "G"), EUR),
           ("2035 revenue", lambda n: link(n, "Total revenue", "rows", "L"), EUR),
           ("2030 EBITDA margin", lambda n: link(n, "EBITDA margin", "rows", "G"), PCT),
           ("Enterprise value", lambda n: link(n, "Enterprise value"), EUR),
           ("Value per share (€)", lambda n: link(n, "Value per share (€)"), EUR2),
           ("Upside / (downside)", lambda n: link(n, "Upside / (downside) vs price"), PCT),
           ("Terminal value as % of EV", lambda n: link(n, "Terminal value as % of EV"), PCT),
           ("Probability", None, PCT)]
for i, (lab, fn, fmt) in enumerate(SUMMARY):
    r = 5 + i
    put(sm, f"A{r}", lab, bold=lab == "Value per share (€)")
    for k, n in enumerate(("Bear", "Base", "Bull")):
        put(sm, f"{'BCD'[k]}{r}", fn(n) if fn else f"=Inputs!B{70 + k}", GREEN, fmt, bold=lab == "Value per share (€)")
put(sm, "E11", "=SUMPRODUCT(B11:D11,B14:D14)", fmt=EUR2, bold=True)
put(sm, "E12", "=E11/Inputs!B5-1", fmt=PCT)
put(sm, "A16", "Current share price (€)")
put(sm, "B16", "=Inputs!B5", GREEN, EUR2)
put(sm, "A17", "WACC")
put(sm, "B17", "=Inputs!B19", GREEN, PCT)
put(sm, "A18", "Orders a year (2027-35) needed to justify the price", bold=True)
put(sm, "B18", "='Reverse DCF'!C7", GREEN, "0.0", bold=True)
put(sm, "A19", "2026 revenue vs company guidance (low end)")
put(sm, "B19", "=C6/Inputs!B76-1", fmt=PCT)
put(sm, "A20", "Dividend yield (FY2025 dividend €8.94)")
put(sm, "B20", "=Historical!F10/Inputs!B5", GREEN, PCT)
put(sm, "A22", "Legend: blue = input, black = formula, green = link to another sheet, yellow = key assumption.", italic=True)
put(sm, "A23", "Sources: GTT press releases FY2025 (2026-02-19) and H1 2026 (2026-07-28); Yahoo Finance 2026-10-16.", italic=True)
put(sm, "A24", "Personal research, not investment advice.", italic=True)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
