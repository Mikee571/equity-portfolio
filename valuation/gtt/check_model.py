"""Prints the key outputs of gtt_dcf.xlsx after it has been recalculated in Excel, and checks for formula errors."""
from openpyxl import load_workbook
wb = load_workbook("gtt_dcf.xlsx", data_only=True); wf = load_workbook("gtt_dcf.xlsx")
n = sum(1 for ws in wf.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
errs = [(ws.title, c.coordinate) for ws in wf.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=") and isinstance(wb[ws.title][c.coordinate].value, str)
        and wb[ws.title][c.coordinate].value.startswith("#")]
print(f"{n} formulas, errors: {errs or 'none'}")
s = wb["Summary"]
for r in list(range(5, 15)) + list(range(16, 21)):
    vals = [s.cell(r, k).value for k in range(2, 6)]
    print(f"{str(s.cell(r,1).value):44s}", *[("%12.1f%%" % (v*100)) if isinstance(v, float) and abs(v) < 1.5 else (f"{v:>12,.1f}" if isinstance(v, (int, float)) else f"{'':>12s}") for v in vals])
d = wb["DCF"]
for top in (3,):
    for lab_row in range(top + 2, top + 25):
        print(f"{str(d.cell(lab_row,1).value):40s}", *[f"{(d.cell(lab_row,k).value or 0):8.1f}" if not isinstance(d.cell(lab_row,k).value, str) else "" for k in range(3, 13)])
