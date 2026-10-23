"""Prints the key outputs of jpm_ri_model.xlsx after it has been recalculated in Excel, and checks for formula errors."""
from openpyxl import load_workbook
wb = load_workbook("jpm_ri_model.xlsx", data_only=True); wf = load_workbook("jpm_ri_model.xlsx")
n = sum(1 for ws in wf.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
errs = [(ws.title, c.coordinate) for ws in wf.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=") and isinstance(wb[ws.title][c.coordinate].value, str)
        and wb[ws.title][c.coordinate].value.startswith("#")]
print(f"{n} formulas, errors: {errs or 'none'}")
s = wb["Summary"]
for r in list(range(5, 12)) + list(range(13, 17)):
    print(f"{str(s.cell(r, 1).value):40s}", [round(v, 4) if isinstance(v, float) else v for v in (s.cell(r, k).value for k in range(2, 6))])
i = wb["Inputs"]; print("shares", round(i["B6"].value, 1), "TCE", round(i["B11"].value), "P/TBV", round(i["B12"].value, 3), "r", round(i["B21"].value, 4))
ri = wb["RI model"]
for row in ri.iter_rows(min_row=1, max_row=ri.max_row):
    if row[0].value and "mechanics" in str(row[0].value):
        for rr in range(row[0].row + 1, row[0].row + 6):
            print(ri.cell(rr, 1).value, [round(ri.cell(rr, k).value, 3) if isinstance(ri.cell(rr, k).value, float) else ri.cell(rr, k).value for k in range(3, 14)])
p = wb["Peers"]
for r in range(5, 11): print([p.cell(r, k).value if k == 1 else round(p.cell(r, k).value, 3) for k in range(1, 9)])
rv = wb["Reverse"]; print("implied ROTE", rv["C6"].value)
for r in range(12, 21): print([rv.cell(r, k).value if not isinstance(rv.cell(r, k).value, float) else round(rv.cell(r, k).value, 3) for k in (1, 3, 4, 5, 6, 7, 8, 9)])
