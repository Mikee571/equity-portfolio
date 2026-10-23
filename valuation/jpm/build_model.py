"""Builds jpm_ri_model.xlsx: a three-scenario residual income (excess return) model of JPMorgan Chase, with live formulas.

Why residual income and not a DCF: for a bank, deposits and debt are raw material rather than financing, so free
cash flow isn't meaningful. Equity value = tangible common equity today + the present value of the profits earned
above the cost of equity on that capital. In steady state this collapses to the justified P/TBV = (ROTE - g) / (r - g).

Sources: JPM Q3 2026 earnings release (8-K, 2026-10-13); SEC XBRL company facts (history, data/jpm_annual_raw.csv);
peers from data/peers.csv (peers.py); prices, beta and 10-year yield from Yahoo Finance at the 2026-10-16 close.
Run: python build_model.py, then open in Excel (or recalculate) to cache values.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).parent
OUT = HERE / "jpm_ri_model.xlsx"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="DDE6F0")
USD = '#,##0;(#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'
I = "Inputs!"
YEARS = 10
YC = [chr(ord("D") + k) for k in range(YEARS)]          # Y1..Y10 in columns D..M; Y0 in column C


def font(color=BLACK, bold=False, italic=False):
    return Font(name="Arial", size=10, color=color, bold=bold, italic=italic)


def put(ws, ref, value, color=BLACK, fmt=None, bold=False, fill=None, italic=False):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def header_row(ws, row, labels, start=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=start + j, value=lab)
        c.font = font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(horizontal="left" if j == 0 else "center")


wb = Workbook()

# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Inputs"
inp.column_dimensions["A"].width = 52
for c in ["B"] + YC + ["C", "N"]:
    inp.column_dimensions[c].width = 9
inp.column_dimensions["O"].width = 90
put(inp, "A1", "JPMorgan Chase: inputs and scenario levers ($mm unless stated)", bold=True)
put(inp, "A2", "Blue = hardcoded input. Yellow = key assumption. Black = formula. Sources in column O.", italic=True)
ROWS = [
    (4, "Market data and capital", None, None, None),
    (5, "Share price ($)", 339.72, USD2, "Yahoo Finance, close 2026-10-16"),
    (6, "Common shares outstanding (mm)", "=(B7-B8)/B9", USD,
     "Derived: (total equity - preferred) / book value per share; Yahoo shows 2,655.6mm"),
    (7, "Total stockholders' equity, 30 Sep 2026", 372000, USD, "Q3 2026 earnings release (8-K exhibit 99.1, 2026-10-13)"),
    (8, "Preferred stock (liquidation preference)", 21200, USD, "SEC XBRL PreferredStockLiquidationPreferenceValue, 30 Jun 2026"),
    (9, "Book value per share ($)", 132.30, USD2, "Q3 2026 earnings release"),
    (10, "Tangible book value per share ($)", 112.64, USD2, "Q3 2026 earnings release"),
    (11, "Tangible common equity (TCE)", "=B10*B6", USD, "Formula"),
    (12, "Current P/TBV", "=B5/B10", MULT, "Formula"),
    (13, "Q3 2026 ROTCE: reported / ex significant items", "20% / 21%", None, "Q3 2026 earnings release ($658mm securities loss)"),
    (14, "CET1 ratio (Standardized)", 0.139, PCT, "Q3 2026 earnings release; ~$30bn excess capital per management"),
    (16, "Cost of equity", None, None, None),
    (17, "Risk-free rate (US 10-year)", 0.051, PCT, "^TNX close 2026-10-16"),
    (18, "Beta (5y monthly)", 1.015, "0.000", "Yahoo Finance"),
    (19, "Adjusted beta (Blume: 0.67 x raw + 0.33)", "=0.67*B18+0.33", "0.000", "Formula"),
    (20, "Equity risk premium", 0.05, PCT, "Assumption (author)"),
    (21, "Cost of equity (r)", "=B17+B19*B20", PCT, "CAPM"),
    (22, "Terminal growth (g)", 0.04, PCT, "Assumption: roughly US nominal GDP; same as the original pitch"),
    (24, "Per-share mechanics (illustrative block in the base case)", None, None, None),
    (25, "Share of distributions done as buybacks", 0.50, PCT, "Assumption; Q3 2026 net buybacks $0.66bn vs dividends $4.4bn"),
    (26, "Buyback price as a multiple of TBVPS", "=B12", MULT, "Assumes today's P/TBV persists"),
]
for r, label, val, fmt, src in ROWS:
    put(inp, f"A{r}", label, bold=val is None)
    if val is None:
        continue
    is_formula = isinstance(val, str) and val.startswith("=")
    put(inp, f"B{r}", val, BLACK if is_formula else BLUE, fmt)
    put(inp, f"O{r}", src, italic=True)
for r in (20, 22):
    inp[f"B{r}"].fill = YELLOW

SCEN = {
    "Bear": dict(start=30, rote=[0.19, 0.12, 0.11, 0.14, 0.15, 0.15, 0.15, 0.15, 0.15, 0.15], lr=0.15, payout=0.45,
                 story="Credit cycle in 2027-28 (card and wholesale losses, reserve builds), then a lower-return normal of 15%."),
    "Base": dict(start=35, rote=[0.20, 0.19, 0.18, 0.175, 0.17, 0.17, 0.17, 0.17, 0.17, 0.17], lr=0.17, payout=0.60,
                 story="Glide from today's ~21% core ROTCE to management's through-the-cycle target of 17%."),
    "Bull": dict(start=40, rote=[0.21, 0.21, 0.21, 0.21, 0.21, 0.205, 0.20, 0.20, 0.20, 0.20], lr=0.20, payout=0.65,
                 story="JPM keeps its edge: ROTE stays above 20% for a decade."),
}
put(inp, "A28", "Scenarios: return on tangible common equity by year (Y1 = Oct 2026 - Sep 2027)", bold=True)
header_row(inp, 29, ["", "", ""] + [f"Y{k + 1}" for k in range(YEARS)])
for name, s in SCEN.items():
    r = s["start"]
    put(inp, f"A{r}", f"{name} case", bold=True)
    put(inp, f"O{r}", s["story"], italic=True)
    put(inp, f"A{r + 1}", "  ROTE")
    for k, col in enumerate(YC):
        put(inp, f"{col}{r + 1}", s["rote"][k], BLUE, PCT)
    put(inp, f"A{r + 2}", "  Long-run ROTE (terminal)")
    put(inp, f"C{r + 2}", s["lr"], BLUE, PCT, fill=YELLOW)
    put(inp, f"A{r + 3}", "  Payout (dividends + buybacks, % of earnings)")
    put(inp, f"C{r + 3}", s["payout"], BLUE, PCT)
put(inp, "O36", "Management's through-the-cycle ROTCE target is 17%; Q3 2026 core ROTCE was 21%.", italic=True)
put(inp, "A45", "Scenario probabilities", bold=True)
for k, (name, p) in enumerate([("Bear", 0.25), ("Base", 0.50), ("Bull", 0.25)]):
    put(inp, f"A{46 + k}", f"  {name}")
    put(inp, f"B{46 + k}", p, BLUE, PCT, fill=YELLOW)
put(inp, "A49", "  Check (must be 100%)")
put(inp, "B49", "=SUM(B46:B48)", fmt=PCT)
put(inp, "A51", "For reference", bold=True)
put(inp, "A52", "  Analyst mean target ($) / number of analysts")
put(inp, "B52", 372.57, BLUE, USD2)
put(inp, "C52", 23, BLUE, "0")
put(inp, "A53", "  Forward EPS, next 12 months ($)")
put(inp, "B53", 25.50, BLUE, USD2)
put(inp, "O52", "Yahoo Finance, 2026-10-16", italic=True)

# ---------------------------------------------------------------- Historical
his = wb.create_sheet("Historical")
raw = pd.read_csv(HERE / "data" / "jpm_annual_raw.csv", index_col=0)
raw.index = raw.index.astype(str)
years = [str(y) for y in range(2012, 2026)]
cols = [chr(ord("B") + i) for i in range(len(years))]
his.column_dimensions["A"].width = 44
for c in cols:
    his.column_dimensions[c].width = 10
put(his, "A1", "JPMorgan Chase: history ($mm unless stated)", bold=True)
put(his, "A2", "Source: SEC EDGAR XBRL company facts (10-K), CIK 0000019617.", italic=True)
header_row(his, 4, ["$mm"] + years)
LINES = [("Net interest income", "NII", USD), ("Noninterest revenue", "NonintRev", USD), ("Noninterest expense", "NonintExp", USD),
         ("Provision for credit losses", "Provision", USD), ("Net income", "NetIncome", USD),
         ("Net income to common", "NI_common", USD), ("Total stockholders' equity", "Equity", USD),
         ("Goodwill", "Goodwill", USD), ("Diluted EPS ($)", "EPS", USD2), ("Dividends declared per share ($)", "DivPerShare", USD2),
         ("Buybacks", "Buybacks", USD), ("Diluted shares (mm)", "DilShares", USD)]
R = {}
for i, (label, key, fmt) in enumerate(LINES):
    r = 5 + i
    R[label] = r
    put(his, f"A{r}", label, bold=label in ("Net income", "Net income to common"))
    for j, y in enumerate(years):
        if y in raw.index and pd.notna(raw.loc[y, key]):
            v = float(raw.loc[y, key])
            v = v if key in ("EPS", "DivPerShare") else v / 1e6
            put(his, f"{cols[j]}{r}", round(v, 2), BLUE, fmt)
r0 = 5 + len(LINES) + 1
put(his, f"A{r0}", "Efficiency ratio (expense / revenue)")
put(his, f"A{r0 + 1}", "Provision / net interest income")
put(his, f"A{r0 + 2}", "Return on total equity (net income / year-end equity)")
put(his, f"A{r0 + 3}", "Dividend payout (DPS / EPS)")
for c in cols:
    put(his, f"{c}{r0}", f"={c}{R['Noninterest expense']}/({c}{R['Net interest income']}+{c}{R['Noninterest revenue']})", fmt=PCT)
    put(his, f"{c}{r0 + 1}", f"={c}{R['Provision for credit losses']}/{c}{R['Net interest income']}", fmt=PCT)
    put(his, f"{c}{r0 + 2}", f"={c}{R['Net income']}/{c}{R['Total stockholders' + chr(39) + ' equity']}", fmt=PCT)
    put(his, f"{c}{r0 + 3}", f"={c}{R['Dividends declared per share ($)']}/{c}{R['Diluted EPS ($)']}", fmt=PCT)
put(his, f"A{r0 + 5}", "2020 shows the credit-cycle pattern: provisions jumped to $17.5bn and net income fell 20%. "
                       "In 2021 reserve releases (negative provision) lifted earnings to a record.", italic=True)

# ---------------------------------------------------------------- Residual income model
ri = wb.create_sheet("RI model")
ri.column_dimensions["A"].width = 44
for c in ["B", "C"] + YC:
    ri.column_dimensions[c].width = 10
put(ri, "A1", "Residual income model by scenario ($mm). Value = TCE today + PV of (net income - r x opening TCE).", bold=True)
FLOW = ["ROTE", "Opening TCE", "Net income to common", "Distributions (dividends + buybacks)", "Closing TCE",
        "Equity charge (r x opening TCE)", "Residual income", "Discount factor", "PV of residual income"]
VAL = ["Long-run ROTE", "Terminal residual income (year 11)", "Terminal value at Y10", "PV of terminal value",
       "Sum of PV of residual income (Y1-Y10)", "Tangible common equity today", "Equity value",
       "Value per share ($)", "Upside / (downside) vs price", "Implied P/TBV", "Terminal value as % of equity value"]
MECH = ["Shares outstanding (mm)", "EPS ($)", "TBVPS ($)", "EPS growth", "TBVPS growth"]
BLOCK = {}


def block(name, top):
    sr = SCEN[name]["start"]
    put(ri, f"A{top}", f"{name} case", bold=True)
    header_row(ri, top + 1, ["$mm", "", "Y0"] + [f"Y{k + 1}" for k in range(YEARS)])
    rr = {lab: top + 2 + i for i, lab in enumerate(FLOW)}
    for lab, r in rr.items():
        put(ri, f"A{r}", lab, bold=lab in ("Residual income", "Net income to common"))
    put(ri, f"C{rr['Closing TCE']}", f"={I}$B$11", GREEN, USD)
    for k, c in enumerate(YC):
        p = "C" if k == 0 else YC[k - 1]
        put(ri, f"{c}{rr['ROTE']}", f"={I}{c}{sr + 1}", GREEN, PCT)
        put(ri, f"{c}{rr['Opening TCE']}", f"={p}{rr['Closing TCE']}", fmt=USD)
        put(ri, f"{c}{rr['Net income to common']}", f"={c}{rr['ROTE']}*{c}{rr['Opening TCE']}", fmt=USD)
        put(ri, f"{c}{rr['Distributions (dividends + buybacks)']}", f"=-{I}$C${sr + 3}*{c}{rr['Net income to common']}", fmt=USD)
        put(ri, f"{c}{rr['Closing TCE']}",
            f"={c}{rr['Opening TCE']}+{c}{rr['Net income to common']}+{c}{rr['Distributions (dividends + buybacks)']}", fmt=USD)
        put(ri, f"{c}{rr['Equity charge (r x opening TCE)']}", f"=-{I}$B$21*{c}{rr['Opening TCE']}", fmt=USD)
        put(ri, f"{c}{rr['Residual income']}",
            f"={c}{rr['Net income to common']}+{c}{rr['Equity charge (r x opening TCE)']}", fmt=USD, bold=True)
        put(ri, f"{c}{rr['Discount factor']}", f"=1/(1+{I}$B$21)^{k + 1}", fmt="0.000")
        put(ri, f"{c}{rr['PV of residual income']}", f"={c}{rr['Residual income']}*{c}{rr['Discount factor']}", fmt=USD)
    last = YC[-1]
    v0 = top + 2 + len(FLOW) + 1
    vr = {lab: v0 + i for i, lab in enumerate(VAL)}
    for lab, r in vr.items():
        put(ri, f"A{r}", lab, bold=lab in ("Equity value", "Value per share ($)"))
    put(ri, f"C{vr['Long-run ROTE']}", f"={I}C{sr + 2}", GREEN, PCT)
    put(ri, f"C{vr['Terminal residual income (year 11)']}",
        f"=(C{vr['Long-run ROTE']}-{I}$B$21)*{last}{rr['Closing TCE']}", fmt=USD)
    put(ri, f"C{vr['Terminal value at Y10']}",
        f"=C{vr['Terminal residual income (year 11)']}/({I}$B$21-{I}$B$22)", fmt=USD)
    put(ri, f"C{vr['PV of terminal value']}", f"=C{vr['Terminal value at Y10']}*{last}{rr['Discount factor']}", fmt=USD)
    put(ri, f"C{vr['Sum of PV of residual income (Y1-Y10)']}", f"=SUM(D{rr['PV of residual income']}:{last}{rr['PV of residual income']})", fmt=USD)
    put(ri, f"C{vr['Tangible common equity today']}", f"={I}$B$11", GREEN, USD)
    put(ri, f"C{vr['Equity value']}", f"=C{vr['Tangible common equity today']}+C{vr['Sum of PV of residual income (Y1-Y10)']}"
                                       f"+C{vr['PV of terminal value']}", fmt=USD, bold=True)
    put(ri, f"C{vr['Value per share ($)']}", f"=C{vr['Equity value']}/{I}$B$6", fmt=USD2, bold=True)
    put(ri, f"C{vr['Upside / (downside) vs price']}", f"=C{vr['Value per share ($)']}/{I}$B$5-1", fmt=PCT)
    put(ri, f"C{vr['Implied P/TBV']}", f"=C{vr['Value per share ($)']}/{I}$B$10", fmt=MULT)
    put(ri, f"C{vr['Terminal value as % of equity value']}", f"=C{vr['PV of terminal value']}/C{vr['Equity value']}", fmt=PCT)
    end = v0 + len(VAL)
    mr = {}
    if name == "Base":
        put(ri, f"A{end + 1}", "Per-share mechanics: what buybacks at ~3x tangible book do to EPS and TBVPS", bold=True)
        mr = {lab: end + 2 + i for i, lab in enumerate(MECH)}
        for lab, r in mr.items():
            put(ri, f"A{r}", f"  {lab}")
        put(ri, f"C{mr['Shares outstanding (mm)']}", f"={I}$B$6", GREEN, USD)
        put(ri, f"C{mr['TBVPS ($)']}", f"={I}$B$10", GREEN, USD2)
        for k, c in enumerate(YC):
            p = "C" if k == 0 else YC[k - 1]
            buyback = f"-{c}{rr['Distributions (dividends + buybacks)']}*{I}$B$25"
            price = f"({I}$B$26*{p}{mr['TBVPS ($)']})"
            put(ri, f"{c}{mr['Shares outstanding (mm)']}", f"={p}{mr['Shares outstanding (mm)']}-{buyback}/{price}", fmt=USD)
            put(ri, f"{c}{mr['EPS ($)']}", f"={c}{rr['Net income to common']}/{p}{mr['Shares outstanding (mm)']}", fmt=USD2)
            put(ri, f"{c}{mr['TBVPS ($)']}", f"={c}{rr['Closing TCE']}/{c}{mr['Shares outstanding (mm)']}", fmt=USD2)
            if k:
                put(ri, f"{c}{mr['EPS growth']}", f"={c}{mr['EPS ($)']}/{p}{mr['EPS ($)']}-1", fmt=PCT)
            put(ri, f"{c}{mr['TBVPS growth']}", f"={c}{mr['TBVPS ($)']}/{p}{mr['TBVPS ($)']}-1", fmt=PCT)
        put(ri, f"A{end + 8}", "Buybacks at ~3x TBV still lift TBVPS compared with paying the same cash as dividends, but they "
                               "dilute it compared with keeping the cash: each $1 spent removes $1 of tangible book and only "
                               "$0.33 of book value's worth of shares.", italic=True)
        end += 9
    BLOCK[name] = dict(rows=rr, vals=vr, mech=mr)
    return end + 2


top = 3
for name in ("Bear", "Base", "Bull"):
    top = block(name, top)

# ---------------------------------------------------------------- Reverse + sensitivity
rev = wb.create_sheet("Reverse")
rev.column_dimensions["A"].width = 50
for c in "BCDEFGHIJ":
    rev.column_dimensions[c].width = 10
put(rev, "A1", "What does the price imply? Steady state: justified P/TBV = (ROTE - g) / (r - g)", bold=True)
put(rev, "A3", "Current P/TBV")
put(rev, "C3", f"={I}B12", GREEN, MULT)
put(rev, "A4", "Cost of equity (r)")
put(rev, "C4", f"={I}B21", GREEN, PCT)
put(rev, "A5", "Terminal growth (g)")
put(rev, "C5", f"={I}B22", GREEN, PCT)
put(rev, "A6", "ROTE the price implies, forever", bold=True)
put(rev, "C6", "=C3*(C4-C5)+C5", fmt=PCT, bold=True)
put(rev, "A7", "  vs management through-the-cycle target")
put(rev, "C7", 0.17, BLUE, PCT)
put(rev, "A8", "  vs Q3 2026 core ROTCE")
put(rev, "C8", 0.21, BLUE, PCT)
put(rev, "A11", "Value per share ($), steady-state formula: cost of equity (rows) x sustained ROTE (columns)", bold=True)
ROTES = [0.13, 0.15, 0.17, 0.19, 0.21, 0.23, 0.25]
COES = [0.08, 0.085, 0.09, 0.095, 0.10, 0.105, 0.11, 0.12]
header_row(rev, 12, ["r \\ ROTE", ""] + [""] * len(ROTES))
for k, x in enumerate(ROTES):
    put(rev, f"{chr(ord('C') + k)}12", x, BLUE, PCT, bold=True)
for i, rr_ in enumerate(COES):
    r = 13 + i
    put(rev, f"A{r}", rr_, BLUE, PCT, bold=True)
    for k in range(len(ROTES)):
        c = chr(ord("C") + k)
        put(rev, f"{c}{r}", f"={I}$B$10*({c}$12-$C$5)/($A{r}-$C$5)", fmt="$#,##0")
put(rev, "A22", "Current price ($)")
put(rev, "C22", f"={I}B5", GREEN, USD2)

# ---------------------------------------------------------------- Peers
pr = wb.create_sheet("Peers")
peers = pd.read_csv(HERE / "data" / "peers.csv")
pr.column_dimensions["A"].width = 12
for c in "BCDEFGHI":
    pr.column_dimensions[c].width = 13
put(pr, "A1", "US large banks, same method for all: TBVPS = book value per share (Yahoo) - (goodwill + intangibles) / shares "
              "(SEC XBRL); ROTE = TTM net income to common / TCE.", bold=True)
put(pr, "A2", "Prices 2026-10-16. JPM's computed TBVPS ($111.96) differs from the reported $112.64 by deferred-tax netting.",
    italic=True)
header_row(pr, 4, ["Bank", "Price", "TBVPS", "ROTE (TTM)", "P/TBV", "P/E (TTM)", "Justified P/TBV*", "Gap"])
for i, row in peers.iterrows():
    r = 5 + i
    put(pr, f"A{r}", row.bank, bold=row.bank == "JPM")
    put(pr, f"B{r}", round(row.price, 2), BLUE, USD2)
    put(pr, f"C{r}", round(row.tbvps, 2), BLUE, USD2)
    put(pr, f"D{r}", round(row.rote, 4), BLUE, PCT)
    put(pr, f"E{r}", f"=B{r}/C{r}", fmt=MULT)
    put(pr, f"F{r}", round(row.p_e, 2), BLUE, '0.0"x"')
    put(pr, f"G{r}", f"=(D{r}-{I}$B$22)/({I}$B$21-{I}$B$22)", fmt=MULT)
    put(pr, f"H{r}", f"=E{r}/G{r}-1", fmt=PCT)
put(pr, "A12", "* If today's TTM ROTE lasted forever, at JPM's cost of equity and growth. Gap > 0 means the price asks for more "
               "than today's returns.", italic=True)

# ---------------------------------------------------------------- Summary
sm = wb.create_sheet("Summary", 0)
sm.column_dimensions["A"].width = 46
for c in "BCDE":
    sm.column_dimensions[c].width = 15
put(sm, "A1", "JPMorgan Chase (JPM): three-scenario residual income model", bold=True)
put(sm, "A2", "Valuation date 2026-10-16. $mm unless stated. Change blue/yellow cells in 'Inputs'; everything else recalculates.",
    italic=True)
header_row(sm, 4, ["", "Bear", "Base", "Bull", "Prob.-weighted"])
SUMMARY = [("ROTE in Y2", lambda n: f"='RI model'!E{BLOCK[n]['rows']['ROTE']}", PCT),
           ("Long-run ROTE", lambda n: f"='RI model'!C{BLOCK[n]['vals']['Long-run ROTE']}", PCT),
           ("Equity value", lambda n: f"='RI model'!C{BLOCK[n]['vals']['Equity value']}", USD),
           ("Value per share ($)", lambda n: f"='RI model'!C{BLOCK[n]['vals']['Value per share ($)']}", USD2),
           ("Upside / (downside)", lambda n: f"='RI model'!C{BLOCK[n]['vals']['Upside / (downside) vs price']}", PCT),
           ("Implied P/TBV", lambda n: f"='RI model'!C{BLOCK[n]['vals']['Implied P/TBV']}", MULT),
           ("Probability", None, PCT)]
for i, (lab, fn, fmt) in enumerate(SUMMARY):
    r = 5 + i
    put(sm, f"A{r}", lab, bold=lab == "Value per share ($)")
    for k, n in enumerate(("Bear", "Base", "Bull")):
        put(sm, f"{'BCD'[k]}{r}", fn(n) if fn else f"=Inputs!B{46 + k}", GREEN, fmt, bold=lab == "Value per share ($)")
put(sm, "E8", "=SUMPRODUCT(B8:D8,B11:D11)", fmt=USD2, bold=True)
put(sm, "E9", "=E8/Inputs!B5-1", fmt=PCT)
put(sm, "A13", "Current price ($) / P/TBV")
put(sm, "B13", "=Inputs!B5", GREEN, USD2)
put(sm, "C13", "=Inputs!B12", GREEN, MULT)
put(sm, "A14", "Cost of equity / terminal growth")
put(sm, "B14", "=Inputs!B21", GREEN, PCT)
put(sm, "C14", "=Inputs!B22", GREEN, PCT)
put(sm, "A15", "ROTE the price implies, forever", bold=True)
put(sm, "B15", "=Reverse!C6", GREEN, PCT, bold=True)
put(sm, "A16", "Analyst mean target ($)")
put(sm, "B16", "=Inputs!B52", GREEN, USD2)
put(sm, "A18", "Legend: blue = input, black = formula, green = link to another sheet, yellow = key assumption.", italic=True)
put(sm, "A19", "Sources: JPM Q3 2026 earnings release (2026-10-13); SEC EDGAR XBRL; Yahoo Finance 2026-10-16.", italic=True)
put(sm, "A20", "Personal research, not investment advice.", italic=True)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
